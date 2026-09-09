"""Excel 课表导入与教学班字段解析。"""

import re
from pathlib import Path


WEEKDAY_NAMES = {"一": "1", "二": "2", "三": "3", "四": "4", "五": "5", "六": "6", "日": "7", "天": "7"}


def _text(value):
    return str(value or "").strip()


def _header_key(value):
    return re.sub(r"\s+", "", _text(value)).lower()


def _find_column(headers, aliases):
    normalized = {_header_key(value): index for index, value in enumerate(headers)}
    for alias in aliases:
        index = normalized.get(_header_key(alias))
        if index is not None:
            return index
    return None


def parse_lesson_time(value):
    """解析“星期四第5-6节{1-16周}”等正方课表时间。"""
    text = _text(value)
    match = re.search(r"(?:星期|周)([一二三四五六日天1-7])[^0-9]{0,8}第?\s*(\d+(?:\s*[-－至到]\s*\d+)?)\s*节", text)
    if not match:
        return "", ""
    weekday = WEEKDAY_NAMES.get(match.group(1), match.group(1))
    section = re.sub(r"\s*[-－至到]\s*", "-", match.group(2))
    return weekday, section


def parse_weeks(value, fallback=""):
    text = _text(value)
    match = re.search(r"\{([^{}]+)\}", text)
    if match:
        text = match.group(1)
    text = text.strip().replace(" ", "")
    if not text:
        text = _text(fallback).strip().replace(" ", "")
    return text.removesuffix("周")


def load_schedule_workbook(path):
    """读取 xlsx/xlsm 课表，返回可用于抢课助手2的教学班记录。"""
    from openpyxl import load_workbook

    workbook_path = Path(path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    records = []
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            try:
                headers = list(next(rows))
            except StopIteration:
                continue

            code_column = _find_column(headers, ("课程代码", "课程号", "课程编码", "kch", "kch_id"))
            title_column = _find_column(headers, ("课程", "课程名称", "课程名", "kcmc"))
            teacher_column = _find_column(headers, ("教师", "教师名称", "任课教师", "jsxm", "jsxx"))
            time_column = _find_column(headers, ("上课时间", "课程时间", "sksj"))
            weeks_column = _find_column(headers, ("上课周次", "起始结束周", "周次", "zcd"))
            class_column = _find_column(headers, ("教学班", "教学班名称", "jxbmc"))
            if code_column is None or title_column is None:
                continue

            for row in rows:
                values = list(row)
                needed = max(code_column, title_column, teacher_column or 0, time_column or 0, weeks_column or 0, class_column or 0)
                if len(values) <= needed:
                    values.extend([None] * (needed + 1 - len(values)))
                course_code = _text(values[code_column])
                title = _text(values[title_column])
                if not course_code or not title:
                    continue
                lesson_time = _text(values[time_column]) if time_column is not None else ""
                weekday, section = parse_lesson_time(lesson_time)
                weeks_value = _text(values[weeks_column]) if weeks_column is not None else ""
                weeks = parse_weeks(lesson_time, weeks_value)
                teacher = _text(values[teacher_column]) if teacher_column is not None else ""
                class_name = _text(values[class_column]) if class_column is not None else ""
                record = {
                    "course_code": course_code,
                    "title": title,
                    "teacher": teacher,
                    "weekday": weekday,
                    "section": section,
                    "weeks": weeks,
                    "time": lesson_time,
                    "class_name": class_name,
                    "source_sheet": worksheet.title,
                }
                if not weekday or not section:
                    # 保留记录便于诊断，但它不会被填入完整的教学班选择。
                    record["parse_warning"] = "无法从上课时间解析星期或节次"
                records.append(record)
    finally:
        workbook.close()

    unique = []
    seen = set()
    for record in records:
        key = tuple(record.get(name, "") for name in ("course_code", "title", "teacher", "weekday", "section", "weeks", "class_name"))
        if key not in seen:
            seen.add(key)
            unique.append(record)
    return unique

